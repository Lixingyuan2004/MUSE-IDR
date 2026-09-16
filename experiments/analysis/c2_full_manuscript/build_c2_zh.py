from __future__ import annotations

import hashlib
import json
import subprocess
import tempfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import build_c2 as base


ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "outputs" / "paper" / "c2_chinese_manuscript"
DOCX_PATH = OUT / "MUSE_IDR_中文论文初稿.docx"
MARKDOWN_PATH = OUT / "MUSE_IDR_中文论文初稿.md"
REPORT_PATH = OUT / "c2_chinese_build_report.json"
FIGURE_DIR = OUT / "figures"

FIGURE_FILES = {
    1: FIGURE_DIR / "figure_1_muse_idr_architecture.png",
    2: FIGURE_DIR / "figure_2_muse_idr_external_performance.png",
    3: FIGURE_DIR / "figure_3_muse_idr_paired_roc_forest.png",
}

FIGURE_SVG_SOURCES = {
    1: base.C1 / "figure_1_locked_a10_architecture.svg",
    2: base.C1 / "figure_2_external_performance.svg",
    3: base.C1 / "figure_3_paired_roc_forest.svg",
}

FIGURE_DIMENSIONS = {
    1: (1500, 690),
    2: (1500, 700),
    3: (1650, 885),
}

FIGURE_TEXT_REPLACEMENTS = {
    1: {
        "Locked A10 cross-architecture ensemble": "MUSE-IDR cross-architecture ensemble",
    },
    2: {
        "A10-locked": "MUSE-IDR",
    },
    3: {
        "A10 minus comparator": "MUSE-IDR minus comparator",
    },
}

TITLE = "MUSE-IDR：一种用于蛋白质内在无序预测的多层多尺度 ESM-2 跨架构集成模型"
RUNNING_TITLE = "MUSE-IDR 内在无序预测模型"
SUBTITLE = "基于 CAID2 和 CAID3 的锁定外部验证及蛋白质水平配对不确定性分析"

ABSTRACT = (
    "蛋白质内在无序预测方法通常在操作性定义和无序残基比例不同的基准数据集上进行比较，"
    "因此笼统的优越性结论往往缺乏稳健性。本研究提出 MUSE-IDR（内部锁定编号 A10-locked），"
    "该模型是一种经过数据泄漏审计的跨架构集成方法，由两个基于冻结 ESM-2-650M 表征的预测分支组成："
    "一支使用末层表征和多尺度卷积上下文，另一支在相同上下文模块之前学习融合最后四层表征。"
    "模型包含两种架构、五个同源隔离交叉验证折和三个随机种子，共 30 个成员，"
    "并在 logit 空间进行固定等权平均。模型在访问 CAID2 或 CAID3 标签之前完成锁定。"
    "在 Disorder-PDB 轨道上，MUSE-IDR 在 CAID2 和 CAID3 上分别取得 0.958/0.934 和 "
    "0.954/0.931 的 ROC-AUC/AUPRC，蛋白质整簇配对 bootstrap 分析显示其 ROC-AUC 分别高于"
    "复现的 LoRA-DR-Suite 650M 基线 0.037 和 0.033。在 Disorder-NOX 轨道上，模型性能较低，"
    "对应指标为 0.781/0.361 和 0.852/0.606，未表现出一致优势。其在四个轨道上的总体表现与"
    "复现的 PUNCH2-Light Released-13 具有统计可比性。锁定后的误差诊断显示，MUSE-IDR 对长无序片段"
    "具有较低错误率，但在低无序比例蛋白质和长有序区段中存在较多假阳性。结果表明，应结合具体评测轨道"
    "理解预测器性能；模型锁定、同源控制和蛋白质水平配对不确定性分析能够提高外部基准比较的可解释性。"
)

KEYWORDS = ["内在无序", "蛋白质语言模型", "ESM-2", "集成学习", "CAID"]


def prepare_muse_idr_figures() -> None:
    """Create manuscript-local figures with the public model name in visible labels."""
    edge = Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe")
    if not edge.is_file():
        raise FileNotFoundError(edge)
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    edge_profile_holder = tempfile.TemporaryDirectory(prefix="muse_idr_edge_")
    edge_profile = Path(edge_profile_holder.name)
    for number, source in FIGURE_SVG_SOURCES.items():
        if not source.is_file():
            raise FileNotFoundError(source)
        svg_text = source.read_text(encoding="utf-8")
        for old, new in FIGURE_TEXT_REPLACEMENTS[number].items():
            if old not in svg_text:
                raise ValueError(f"figure {number} is missing expected text: {old}")
            svg_text = svg_text.replace(old, new)
        svg_path = FIGURE_DIR / f"figure_{number}_muse_idr.svg"
        svg_path.write_text(svg_text, encoding="utf-8", newline="\n")
        width, height = FIGURE_DIMENSIONS[number]
        output = FIGURE_FILES[number]
        command = [
            str(edge),
            "--headless",
            "--disable-gpu",
            "--hide-scrollbars",
            "--force-device-scale-factor=1",
            f"--user-data-dir={edge_profile}",
            f"--window-size={width},{height}",
            f"--screenshot={output}",
            svg_path.resolve().as_uri(),
        ]
        subprocess.run(command, check=True, capture_output=True, text=True, timeout=120)
        if not output.is_file():
            raise RuntimeError(f"Edge did not render figure {number}: {output}")


def h(level: int, text: str) -> dict:
    return {"type": f"h{level}", "text": text}


def p(text: str) -> dict:
    return {"type": "p", "text": text}


def bullets(items: list[str]) -> dict:
    return {"type": "bullets", "items": items}


def figure(number: int) -> dict:
    return {"type": "figure", "number": number}


def table(number: int) -> dict:
    return {"type": "table", "number": number}


CONTENT = [
    h(1, "1 引言"),
    p(
        "内在无序蛋白质及内在无序区在生理条件下不形成单一、稳定的三级结构，却广泛参与细胞信号传导、"
        "调控、分子识别和大分子组装（Wright and Dyson, 2015）。其构象异质性和环境依赖性增加了实验表征"
        "难度，也使基于序列的计算预测成为重要补充。DisProt 提供人工整理并具有实验支持的无序注释，"
        "蛋白质内在无序预测关键评估（Critical Assessment of protein Intrinsic Disorder prediction, CAID）"
        "则使用此前未见的蛋白质评测不同方法（Piovesan et al., 2017; Necci et al., 2021）。"
    ),
    p(
        "近年的 CAID 评测既显示了方法性能的快速进步，也揭示了持续存在的基准依赖性。Disorder-NOX 参考集"
        "将缺少合格无序证据的残基视为有序，而 Disorder-PDB 会屏蔽缺少已观测结构坐标的负类残基。"
        "因此，两条轨道的类别比例以及可靠负证据的定义均不相同（Del Conte et al., 2023; Mehdiabadi et al., 2026）。"
        "同一方法可能在一条轨道上排名靠前、在另一条轨道上表现一般，而两种结果均可能真实反映相应评测定义。"
        "若只报告单一合并分数，可能掩盖这种操作性异质性。"
    ),
    p(
        "蛋白质语言模型（protein language model, PLM）能够从大规模序列集合中学习具有上下文依赖性的残基表征。"
        "ESM-2 表征编码了进化和结构规律，已成为无序预测器常用的输入特征（Lin et al., 2023; Mehdiabadi et al., 2026）。"
        "现有方法以不同方式利用这些表征：LoRA-DR-Suite 通过低秩更新适配 PLM 的注意力投影，PUNCH2-Light 则通过"
        "深度卷积集成结合 ProtT5 与 one-hot 特征（Lombardi et al., 2025; Meng and Pollastri, 2025）。"
        "这促成了一个互补问题：在严格控制数据泄漏的条件下，仅依靠完全冻结的骨干网络，并结合表征层次和局部上下文"
        "的多样性，能够达到怎样的预测性能？"
    ),
    p(
        "为回答这一问题，我们构建了 MUSE-IDR（内部锁定编号 A10-locked），它是在冻结 ESM-2-650M 残基表征之上"
        "训练的 30 成员跨架构集成模型。第一类成员使用最后一层表征和并行多尺度卷积预测头；第二类成员先学习"
        "第 30-33 层的标量混合，再进入相同的上下文预测模块。所有交叉验证折、随机种子、架构选择和聚合权重均在访问"
        "CAID2 或 CAID3 标签之前固定。本研究的目标为：（i）量化 MUSE-IDR 在 CAID2 和 CAID3 的 Disorder-NOX 与"
        "Disorder-PDB 轨道上的外部性能；（ii）与本地复现且逐残基对齐的 LoRA-DR-Suite 和 PUNCH2-Light 基线进行比较；"
        "（iii）识别能够解释轨道特异性表现的蛋白质和片段类型。"
    ),

    h(1, "2 材料与方法"),
    h(2, "2.1 研究设计与模型锁定"),
    p(
        "研究流程分为开发阶段和锁定后的外部评测阶段。在读取外部标签之前，我们已确定模型架构、训练数据、"
        "交叉验证划分、随机种子、集成成员、成员权重、概率阈值和输出定义。锁定模型包含 30 个检查点哈希值："
        "来自五个交叉验证折以及随机种子 17、29 和 43 的 15 个 A2 成员和 15 个 A8 成员。每个成员以 1/30 的"
        "权重贡献一个 logit。未使用折外标签或 CAID 标签拟合集成权重、重新校准概率或选择轨道特异性变体。"
    ),
    h(2, "2.2 训练数据与数据泄漏控制"),
    p(
        "训练语料包含 1,133 条蛋白质和 573,421 个残基，其中 202,662 个残基具有已知二分类标签，包括 105,036 个"
        "无序残基和 97,626 个有序残基；其余残基均标记为未知，并从损失函数和评价指标中排除。正类标签来源于历史"
        "DisProt 2018_11 快照中的直接经典无序证据。负类标签要求候选序列能够映射至蛋白质结构数据库（PDB）中的"
        "实验观测坐标。X 射线结构的分辨率限制为不高于 3.0 Å，冷冻电镜结构限制为不高于 4.0 Å，核磁共振结构不设"
        "分辨率阈值。已注释为无序、熔融球或前熔融球的区域及其两侧五个残基的边界缓冲区均不得作为负类。"
    ),
    p(
        "在交叉验证之前，我们从标识符、完全相同序列和同源性三个层面排除与 CAID2/CAID3 的潜在重叠。"
        "标识符审计覆盖 DisProt、UniProt 及异构体根标识符；精确序列审计使用标准化序列 SHA256；同源性审计在"
        "MMseqs2 报告序列一致性大于 30%，且查询与目标双方覆盖率均至少为 80% 时排除候选蛋白质。分析使用"
        "MMseqs2 18-8cc5c，并保留原始搜索输出及其哈希值（Steinegger and Soding, 2017）。排除完成后，以相同标准"
        "构建训练集内部全对全同源图。图中 984 个连通分量（190 条非自身边）被完整分配到五个折，同时平衡正类残基、"
        "负类残基、蛋白质数量和总长度。折划分使用随机种子 17，并在所有模型初始化种子之间复用；不存在跨折的合格同源边。"
    ),
    h(2, "2.3 冻结的 ESM-2 表征"),
    p(
        "所有模型均使用 facebook/esm2_t33_650M_UR50D，固定版本为"
        " 08e4846e537177426273712802403f7ba8261b6c，且骨干参数全部冻结。对于超过 1,022 个内容残基的序列，"
        "采用步长为 511 的重叠窗口处理，并通过确定性的中心加权平均融合重叠区域的 logit 或表征。A2 缓存最后一层"
        "形状为 [L,1280] 的表征；A8 独立生成并校验第 30、31、32 和 33 层的 float16 缓存，形状为 [L,4,1280]。"
        "所有缓存均不包含残基标签。"
    ),
    h(2, "2.4 A2 与 A8 预测头"),
    p(
        "A2 首先对每个 1,280 维残基向量进行 LayerNorm，并通过 GELU 投影到 256 维。随后设置四个深度可分离"
        "一维卷积分支，卷积核大小分别为 3、7、15 和 31。各分支输出通过可学习门控和残差连接融合，经过 0.2 的"
        "dropout 后输出单个残基 logit。该设计在不更新 ESM-2 的前提下显式引入多个局部上下文尺度。"
    ),
    p(
        "A8 使用与 A2 相同的上下文预测头，但在此之前先计算 ESM-2 第 30-33 层的全局可学习 softmax 混合。"
        "混合权重初始化为 [0.1, 0.1, 0.1, 0.7]，使初始状态偏向最后一层。与 A2 相比，A8 仅增加四个可训练标量参数。"
        "两类模型都只有一个输出头，用于表示经典内在无序概率；均不预测软无序，也不产生针对 NOX 或 PDB 的轨道特异性输出。"
    ),
    figure(1),
    h(2, "2.5 模型训练与集成构建"),
    p(
        "预测头最多训练 20 个 epoch，优化器为 AdamW，学习率为 3×10^-4，权重衰减为 0.01，蛋白质批大小为 4，"
        "梯度范数裁剪阈值为 1.0。A8 的层混合参数单独使用 3×10^-3 的学习率。带 logit 的二元交叉熵仅在已知标签"
        "残基上计算，并使用各折负类与正类残基数量之比作为正类权重。验证集 micro ROC-AUC 用于检查点选择和早停，"
        "早停耐心值为 4。外部推理时，先对 30 个成员的 logit 进行固定等权平均，再通过 sigmoid 函数转换为概率。"
    ),
    h(2, "2.6 外部数据集与复现基线"),
    p(
        "外部评测使用 CAID2 和 CAID3 的 Disorder-NOX 与 Disorder-PDB 参考集。CAID2-NOX 包含 210 条蛋白质和"
        "160,802 个已知标签残基，CAID2-PDB 包含 348 条蛋白质和 130,877 个已知标签残基；CAID3-NOX 包含 204 条"
        "蛋白质和 99,977 个已知标签残基，CAID3-PDB 包含 319 条蛋白质和 99,239 个已知标签残基。所有被屏蔽标签均"
        "从评测中排除。MUSE-IDR 的预测仅从不含标签的 FASTA 文件生成一次，完成哈希锁定后再与各参考集连接。"
    ),
    p(
        "我们复现了三个符合条件的基线变体。LoRA-DR-Suite 使用作者发布的 ESM-2-650M 适配器，并仅选择在 DisProt 7"
        "上训练的版本；使用过 CAID 标签的变体和所有软无序输出头均被排除。推理前校验发布的适配器、保存的 token 分类头"
        "以及 query/value LoRA 张量。PUNCH2-Light 按提交版本 6c7935b3597c056d2e6b3845bb54fc101c5bc574 复现，"
        "并使用官方 one-hot 与 ProtT5 特征。由于论文描述包含 8 个集成成员，而发布入口实际载入 13 个成员，"
        "因此分别报告 Paper-8 和 Released-13 两个变体。所有比较模型的预测均与相同参考集逐残基对齐，并在评价前完成哈希锁定。"
    ),
    h(2, "2.7 评价指标与验证性统计"),
    p(
        "我们报告全精度 ROC-AUC、梯形积分得到的精确率-召回率曲线下面积（AUPRC）、平均精确率（APS）、预先规定阈值"
        "0.5 下的 F1 和 Matthews 相关系数（MCC）、Fmax 以及 Brier 分数。由于不同评测轨道的类别不平衡程度不同，"
        "同时保留精确率-召回率指标（Davis and Goadrich, 2006; Saito and Rehmsmeier, 2015）。bootstrap 重采样以完整"
        "蛋白质为单位，从而保留同一蛋白质内部残基之间的依赖关系。"
    ),
    p(
        "在每条评测轨道上，四个模型共享 2,000 次配对 bootstrap 重采样。计算 MUSE-IDR 减去比较模型所得指标差值的"
        "百分位数 95% 置信区间和双侧 bootstrap p 值，并对 ROC-AUC 额外执行配对 DeLong 检验（DeLong et al., 1988）。"
        "Holm 校正控制 48 个主要假设检验，覆盖 MUSE-IDR 与 LoRA-DR-Suite、MUSE-IDR 与 PUNCH2-Light Released-13"
        "在四条轨道和六项指标上的比较。Paper-8 比较构成单独的 24 项补充检验族（Holm, 1979）。"
    ),
    h(2, "2.8 锁定后的描述性分析"),
    p(
        "误差和互补性分析均属于描述性分析，不增加新的验证性假设检验。预先规定的蛋白质长度分层为 1-200、201-500、"
        "501-1000 和 1001 个残基以上；已知标签中的无序比例分层为 0-0.10、(0.10,0.30]、(0.30,0.60] 和 (0.60,1.00]。"
        "残基层面的诊断变量包括到最近 0/1 标签转换点的距离，以及真实有序或无序连续区段长度（1-15、16-30、31-100 和"
        "101 个残基以上）。使用阈值错误分歧、仅单一模型预测正确的比例以及 Spearman 分数相关系数描述模型互补性。"
        "这些结果均未用于修改 MUSE-IDR。"
    ),
    h(2, "2.9 可重复性与软件环境"),
    p(
        "所有源代码压缩包、模型检查点、预测文件、外部参考集和分析输出均由 SHA256 清单保护。锁定推理烟雾测试要求"
        "完整载入 30 个成员、排除开发检查点、验证每个输入残基恰有一个预测概率，并确认缓存复用前后的预测完全一致。"
        "在 CAID3 两条轨道上，内部 CAID 兼容指标与官方 CAID 实现版本"
        " 7bab7e8880d7f949e480fdb377c5a5c8546ae336 完全一致。分析环境包括 Python 3.12、PyTorch 2.12.1、"
        "Transformers 5.15.1 和 MMseqs2 18-8cc5c。训练和推理均在 NVIDIA GeForce RTX 4090 上完成。"
    ),

    h(1, "3 结果"),
    h(2, "3.1 锁定集成模型的外部性能"),
    p(
        "MUSE-IDR 通过组合两种表征视角获得模型多样性，而不是为不同评测轨道拟合更大的专用预测头。15 个 A2 成员"
        "使用 ESM-2 最后一层表征，15 个 A8 成员学习最后四层的混合。固定等权 logit 集成在 CAID2-PDB 和 CAID3-PDB"
        "上分别取得 0.958/0.934 和 0.954/0.931 的 ROC-AUC/AUPRC。NOX 轨道上的表现较低且具有数据集依赖性："
        "CAID2 为 0.781/0.361，CAID3 为 0.852/0.606（表 1；图 2）。"
    ),
    table(1),
    figure(2),
    h(2, "3.2 验证性配对比较"),
    p(
        "与 LoRA-DR-Suite 650M 相比，MUSE-IDR 在两条 PDB 轨道上的 ROC-AUC 更高。CAID2-PDB 上的配对差值为"
        "+0.0374（蛋白质 bootstrap 95% CI：+0.0227 至 +0.0534；Holm 校正 bootstrap p=0.048），CAID3-PDB"
        "上的差值为 +0.0334（+0.0192 至 +0.0484；校正 p=0.048）。在 CAID2-NOX 上，MUSE-IDR 低 0.0541"
        "（-0.0973 至 -0.0160），但在 48 项检验的 Holm 校正后 bootstrap p 值为 0.180。CAID3-NOX 的差值为"
        "-0.0079，置信区间跨越 0。"
    ),
    p(
        "在四条轨道上，MUSE-IDR 与 PUNCH2-Light Released-13 的蛋白质配对 bootstrap 区间均有重叠。"
        "MUSE-IDR 的点估计在 CAID2-NOX、CAID2-PDB、CAID3-NOX 和 CAID3-PDB 上分别高 0.0013、0.0078、"
        "0.0122 和 0.0011。这些结果支持二者具有可比的外部性能，而不支持 MUSE-IDR 在所有条件下普遍优越。"
        "补充的 Paper-8 比较方向相近，但应在其独立的多重校正检验族中解释（表 2；图 3）。"
    ),
    table(2),
    figure(3),
    h(2, "3.3 无序比例与片段长度相关的性能区间"),
    p(
        "锁定后的分层分析将主要弱点定位于已知无序残基比例为 0-10% 的蛋白质。在该分层中，CAID2-NOX 和"
        "CAID3-NOX 上 MUSE-IDR 减去 LoRA-DR-Suite 的 ROC-AUC 差值约为 -0.090 和 -0.058。随着无序比例升高，"
        "差值方向发生逆转：在 CAID2-NOX 的 30-60% 分层中约为 +0.105，在 CAID3-NOX 的大于 60% 分层中约为"
        "+0.085。这些探索性结果见补充图 S1 和补充表 S2。"
    ),
    p(
        "对于长度至少为 101 个残基的真实无序连续区段，MUSE-IDR 在四条轨道上固定阈值错误率约为 0.047-0.072；"
        "LoRA-DR-Suite 和 PUNCH2-Light Released-13 的对应范围分别为 0.226-0.311 和 0.120-0.133。相反，"
        "MUSE-IDR 在 NOX 的长有序区段中产生了更多假阳性。在 CAID2-NOX 上，尽管无序标签比例为 0.195，模型平均"
        "预测概率达到 0.515，提示存在正向校准偏差（补充图 S2；补充表 S3-S4）。"
    ),
    h(2, "3.4 残基层面的预测互补性"),
    p(
        "MUSE-IDR 与复现基线并不会在完全相同的位置出错。相对于 LoRA-DR-Suite，MUSE-IDR 在 CAID2-PDB 和"
        "CAID3-PDB 上分别有 8.8% 和 7.8% 的残基仅由自身预测正确，而仅比较模型预测正确的比例为 2.6% 和 4.3%。"
        "在 NOX 轨道上方向相反，LoRA 独有正确率高于 MUSE-IDR。MUSE-IDR 与 PUNCH2-Light 的预测分数高度相关"
        "（Spearman 约为 0.83-0.90），但根据轨道和模型变体不同，二者阈值化结果仍在 5.6%-11.8% 的残基上不一致"
        "（补充表 S5）。"
    ),

    h(1, "4 讨论"),
    h(2, "4.1 主要发现"),
    p(
        "MUSE-IDR 最适合被理解为一次经过数据泄漏审计、重视评测轨道差异的冻结 PLM 跨架构集成模型外部验证。"
        "该模型在 Disorder-PDB 和长无序片段上表现突出；与复现的 LoRA-DR-Suite 相比，其 PDB 优势在完整蛋白质"
        "配对重采样和家族错误率校正后仍得到支持。同时，MUSE-IDR 并未在 Disorder-NOX 上占据优势，并且在四个主要"
        "数据集-轨道组合中与 PUNCH2-Light Released-13 具有统计可比性。与笼统宣称达到普遍最优相比，这种有边界的"
        "结论更稳健，也更容易复现。"
    ),
    h(2, "4.2 评测轨道定义为何重要"),
    p(
        "NOX 与 PDB 的差异是理解结果的关键。Disorder-NOX 会将注释无序区之外的高分预测视为错误，即使这些位置缺少"
        "完整的结构性负证据；Disorder-PDB 则屏蔽大量此类残基，重点评价实验支持的正类和已观测的有序负类。"
        "因此，MUSE-IDR 对局部呈现无序特征的有序区域持续给出较高分数，会降低其在低无序比例 NOX 蛋白质上的表现，"
        "却有助于完整保留长无序区段。同一种预测行为在不同操作性参考集下可能有利，也可能不利。CAID 的多轨道设计"
        "因此提供了互补信息，而不是重复评价。"
    ),
    h(2, "4.3 架构解释"),
    p(
        "MUSE-IDR 集成了两种计算成本较低的多样性来源。A2 从 ESM-2 最后一层表征中显式融合四种尺度的局部上下文；"
        "A8 在相同预测头之前学习最后四层表征的混合，仅增加四个可训练标量。进一步在不同架构、同源隔离折和随机初始化"
        "之间进行平均，可降低预测对单一表征视角或单次数据划分的依赖。长无序区段上的低错误率与多尺度上下文聚合的设计"
        "相一致，但外部评测本身不能证明因果关系；如需做出组件层面的因果结论，仍需要在重新锁定的开发协议下进行消融实验。"
    ),
    h(2, "4.4 与复现基线的关系"),
    p(
        "LoRA-DR-Suite 会适配 PLM 骨干，并在 NOX 上表现更强，提示任务特异性的注意力更新可能有助于低无序比例条件下的"
        "判别或校准。MUSE-IDR 则保持骨干完全冻结，并在两条 PDB 轨道上更强。PUNCH2-Light 使用 ProtT5 和 one-hot"
        "输入，其总体性能与 MUSE-IDR 接近，但残余预测分歧说明不同表征家族仍具有互补性。未来可探索不使用外部标签的"
        "堆叠、门控或蒸馏，但任何发生改变的预测器都必须使用新的模型标识符，并在未使用过的外部数据上重新评价，"
        "不能作为 MUSE-IDR 锁定版本的事后修订。"
    ),
    h(2, "4.5 局限性"),
    bullets([
        "外部证据仅覆盖 CAID2 和 CAID3 的经典无序轨道，不能代表结合型无序、连接区或软无序任务上的性能。",
        "NOX 和 PDB 是不同的操作性参考定义，任何一方都不能被视为普遍完整的生物学真值。",
        "30 成员集成可以共享骨干表征提取，但其存储和推理开销仍高于单一预测头。",
        "固定阈值和 Brier 分析显示模型在低无序比例 NOX 数据上存在过度预测；本研究有意禁止使用外部标签重新校准。",
        "验证性比较仅覆盖本地复现的 LoRA-DR-Suite 和 PUNCH2-Light 变体；官方排名背景不能等同于逐残基对齐的本地复现。",
        "亚组、边界、片段和互补性分析均为描述性结果，不应解释为新的验证性假设检验。",
    ]),
    h(2, "4.6 未来工作"),
    p(
        "后续模型应仅使用训练集交叉验证针对已识别的失败类型进行改进。候选策略包括增加具有实验支持的有序负类采样、"
        "加入校准正则化、使用边界感知目标，或设计轻量门控以抑制低无序蛋白质上的整体过度预测。模型蒸馏可降低 30 个"
        "预测头带来的成本。MUSE-IDR 与 LoRA-DR-Suite、PUNCH2-Light 的互补性也支持预先规定的跨家族集成研究，"
        "但其权重必须在不使用 CAID2/CAID3 标签的条件下学习，且外部测试集必须保持真正未使用状态。"
    ),

    h(1, "5 结论"),
    p(
        "MUSE-IDR 在严格的同源控制、哈希校验和模型锁定协议下，将末层多尺度 ESM-2 预测头与可学习多层 ESM-2 预测头"
        "进行集成。该模型在 CAID2/CAID3 Disorder-PDB 和长无序片段上表现突出，与复现的 PUNCH2-Light 具有可比性，"
        "同时在低无序比例 NOX 蛋白质上暴露出明确弱点。因此，本研究的贡献既包括模型，也包括评价规范：外部标签仅用于"
        "锁定后的评测，不确定性分析以蛋白质为配对单位，最终结论服从各评测基准的具体定义，而不是被压缩为单一的普遍优越性主张。"
    ),
    h(1, "数据可用性"),
    p(
        "CAID 参考数据和预测格式规范可从 CAID 项目获得，DisProt 和 PDB 源数据可从相应公共资源获得。本研究所使用的"
        "历史快照、派生训练清单、模型检查点和结果压缩包均由项目中的 SHA256 清单唯一标识。"
        "[投稿前插入代码仓库 DOI/URL]。对于许可证缺失或不明确的第三方权重，本研究不进行重新分发。"
    ),
    h(1, "代码可用性"),
    p(
        "训练、锁定推理、基线复现、指标计算、配对 bootstrap 和论文材料生成脚本计划发布于"
        " [代码仓库 URL/DOI]。正式发布内容应包括环境配置、源代码版本锁和最小化的纯序列推理示例。"
    ),
    h(1, "致谢"),
    p("[由作者补充致谢内容]。"),
    h(1, "经费支持"),
    p("[由作者补充资助来源和项目编号]。"),
    h(1, "利益冲突"),
    p("作者声明：[无利益冲突／由作者补充具体冲突]。"),
    h(1, "作者贡献"),
    p("[作者名单确定后，根据 CRediT 贡献角色补充]。"),
    h(1, "人工智能辅助写作声明"),
    p(
        "OpenAI Codex 被用于辅助整理材料、起草语言和以程序方式组装 Word 文档。它未参与模型选择，未在模型开发阶段"
        "访问评测标签，也未解释未经人工审核的原始实验数据。所有科学陈述、引文、分析和最终措辞均须由人类作者核验并批准。"
    ),
]


FIGURE_LEGENDS = {
    1: (
        "MUSE-IDR 的锁定架构与评价流程。冻结的 ESM-2-650M 残基表征进入两个模型家族：A2 使用最后一层表征和"
        "多尺度上下文预测头；A8 在上下文预测前学习第 30-33 层表征的混合。每个家族包含三个随机种子和五个交叉验证折，"
        "共得到 30 个 logit，并通过固定等权 logit 平均进行融合。所有模型选择均在访问 CAID2/CAID3 标签之前完成。"
    ),
    2: (
        "MUSE-IDR 与三个复现基线变体在 CAID2 和 CAID3 外部测试集上的全精度 ROC-AUC 与 AUPRC。"
        "所有模型在对齐比较中均具有完整的残基覆盖率。"
    ),
    3: (
        "ROC-AUC 的蛋白质配对 bootstrap 差值。点表示 MUSE-IDR 减去比较模型，横线表示 2,000 次完整蛋白质"
        "配对 bootstrap 得到的百分位数 95% 置信区间；星号表示 Holm 校正 bootstrap p<0.05。Paper-8 比较"
        "属于单独的补充多重校正检验族。"
    ),
}


def sha256_file(path: Path) -> str:
    hsh = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            hsh.update(chunk)
    return hsh.hexdigest()


def set_run_font_zh(run, name: str = "Calibri", size: float | None = None,
                    color: str | None = None, bold: bool | None = None,
                    italic: bool | None = None) -> None:
    base._ORIGINAL_SET_RUN_FONT(run, name=name, size=size, color=color, bold=bold, italic=italic)
    rfonts = run._element.get_or_add_rPr().rFonts
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def configure_styles_zh(doc: Document) -> None:
    base._ORIGINAL_CONFIGURE_STYLES(doc)
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Caption", "List Bullet", "List Number"):
        style = doc.styles[style_name]
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.45
    doc.styles["Normal"].paragraph_format.space_after = Pt(7)


def configure_page_zh(doc: Document) -> None:
    base._ORIGINAL_CONFIGURE_PAGE(doc)
    for section in doc.sections:
        header = section.header
        hp = header.paragraphs[0]
        hp.clear()
        hp.paragraph_format.space_after = Pt(0)
        hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
        left = hp.add_run(RUNNING_TITLE)
        set_run_font_zh(left, size=8.5, color="6B7280")
        right = hp.add_run("\t原创研究 | 中文初稿")
        set_run_font_zh(right, size=8.5, color="6B7280")


def add_title_page_zh(doc: Document) -> None:
    override = base.DESIGN["journal_title_page_compact"]
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(override["top_spacer_pt"])

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(12)
    set_run_font_zh(kicker.add_run("原创研究 | 中文作者审阅稿"), size=9.5,
                    color=override["kicker_color"], bold=True)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(10)
    title_p.paragraph_format.keep_with_next = True
    set_run_font_zh(title_p.add_run(TITLE), size=20, color=override["title_color"], bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(22)
    set_run_font_zh(subtitle.add_run(SUBTITLE), size=12.5,
                    color=override["subtitle_color"], italic=True)

    for text, size, bold in (
        ("[作者姓名与 ORCID]", 11.5, True),
        ("[作者单位]", 10.5, False),
        ("[通讯作者地址与电子邮箱]", 10.0, False),
    ):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(5)
        set_run_font_zh(para.add_run(text), size=size, color="374151", bold=bold)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(22)
    meta.paragraph_format.space_after = Pt(4)
    set_run_font_zh(meta.add_run("目标期刊：Bioinformatics Advances"), size=10,
                    color="203748", bold=True)

    meta2 = doc.add_paragraph()
    meta2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font_zh(meta2.add_run(f"短标题：{RUNNING_TITLE} | 初稿日期：{date.today().isoformat()}"),
                    size=9.5, color="6B7280")
    doc.add_page_break()


def add_abstract_zh(doc: Document) -> None:
    doc.add_heading("摘要", level=1)
    para = doc.add_paragraph(ABSTRACT)
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para.paragraph_format.keep_together = True
    key = doc.add_paragraph()
    key.paragraph_format.space_before = Pt(3)
    label = key.add_run("关键词：")
    set_run_font_zh(label, bold=True)
    set_run_font_zh(key.add_run("；".join(KEYWORDS)))


def load_table_zh(number: int) -> list[list[str]]:
    rows = base._ORIGINAL_LOAD_TABLE(number)
    if number == 1:
        rows[0] = ["数据集", "轨道", "模型", "ROC-AUC", "AUPRC", "APS", "F1@0.5", "MCC@0.5"]
    else:
        rows[0] = ["数据集", "轨道", "比较模型", "蛋白质数", "MUSE-IDR 差值", "95% CI", "Bootstrap Holm p", "DeLong Holm p"]
    for row in rows[1:]:
        row[1] = row[1].replace("disorder_nox", "NOX").replace("disorder_pdb", "PDB")
        row[2] = row[2].replace("A10-locked", "MUSE-IDR")
    return rows


def add_figure_zh(doc: Document, number: int) -> None:
    image_path = FIGURE_FILES[number]
    if not image_path.is_file():
        raise FileNotFoundError(image_path)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(8)
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.keep_with_next = True
    run = para.add_run()
    run.add_picture(str(image_path), width=Inches(6.15))
    shape = doc.inline_shapes[-1]
    shape._inline.docPr.set("title", f"图 {number}")
    shape._inline.docPr.set("descr", FIGURE_LEGENDS[number])
    caption = doc.add_paragraph(style="Caption")
    caption.paragraph_format.keep_together = True
    lead = caption.add_run(f"图 {number}．")
    set_run_font_zh(lead, bold=True)
    set_run_font_zh(caption.add_run(FIGURE_LEGENDS[number]))


def add_table_zh(doc: Document, number: int) -> None:
    rows = load_table_zh(number)
    if number == 1:
        caption_text = "外部测试集的全精度性能。F1 和 MCC 均使用固定阈值 0.5。"
        widths = [900, 900, 2450, 950, 1000, 950, 950, 1260]
    else:
        caption_text = "基于 2,000 次完整蛋白质 bootstrap 的配对 ROC-AUC 差值（MUSE-IDR 减去比较模型）。"
        widths = [800, 900, 2050, 1050, 1850, 900, 910, 900]

    cap = doc.add_paragraph(style="Caption")
    cap.paragraph_format.space_before = Pt(8)
    cap.paragraph_format.space_after = Pt(4)
    cap.paragraph_format.keep_with_next = True
    lead = cap.add_run(f"表 {number}．")
    set_run_font_zh(lead, bold=True)
    set_run_font_zh(cap.add_run(caption_text))

    tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.style = "Table Grid"
    base.apply_table_geometry(tbl, widths)
    base.set_repeat_table_header(tbl.rows[0])
    for table_row in tbl.rows:
        base.set_table_row_cant_split(table_row)
    for ridx, row in enumerate(rows):
        for cidx, value in enumerate(row):
            cell = tbl.cell(ridx, cidx)
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx == 0:
                base.shade_cell(cell, base.DESIGN["tables"]["header_fill"])
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.space_after = Pt(1.5)
                para.paragraph_format.line_spacing = 1.0
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if cidx in (0, 1, 2, 4) else WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    set_run_font_zh(run, size=7.5 if number == 2 else 7.8,
                                    bold=(ridx == 0), color="111827")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def render_markdown() -> str:
    lines = [
        f"# {TITLE}", "",
        "**文章类型：** 原创研究  ",
        "**目标期刊：** Bioinformatics Advances  ",
        f"**短标题：** {RUNNING_TITLE}", "",
        "**作者：** [作者姓名与 ORCID]  ",
        "**单位：** [作者单位]  ",
        "**通讯作者：** [通讯作者地址与电子邮箱]", "",
        "## 摘要", "", ABSTRACT, "",
        "**关键词：** " + "；".join(KEYWORDS), "",
    ]
    for item in CONTENT:
        kind = item["type"]
        if kind.startswith("h"):
            lines.extend(["#" * (int(kind[1:]) + 1) + " " + item["text"], ""])
        elif kind == "p":
            lines.extend([item["text"], ""])
        elif kind == "bullets":
            lines.extend([f"- {entry}" for entry in item["items"]])
            lines.append("")
        elif kind == "figure":
            number = item["number"]
            lines.extend([
                f"![图 {number}]({FIGURE_FILES[number].as_posix()})", "",
                f"**图 {number}．** {FIGURE_LEGENDS[number]}", "",
            ])
        elif kind == "table":
            number = item["number"]
            rows = load_table_zh(number)
            lines.append("| " + " | ".join(rows[0]) + " |")
            lines.append("| " + " | ".join(["---"] * len(rows[0])) + " |")
            lines.extend(["| " + " | ".join(row) + " |" for row in rows[1:]])
            lines.append("")
    lines.extend(["## 参考文献", ""])
    lines.extend(base.REFERENCES)
    lines.append("")
    return "\n".join(lines)


def build_docx(path: Path) -> None:
    doc = Document()
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "MUSE-IDR 蛋白质内在无序预测模型的锁定外部验证"
    doc.core_properties.author = "匿名作者组"
    doc.core_properties.last_modified_by = "匿名作者组"
    doc.core_properties.keywords = ", ".join(KEYWORDS)
    configure_styles_zh(doc)
    configure_page_zh(doc)
    add_title_page_zh(doc)
    add_abstract_zh(doc)

    for item in CONTENT:
        kind = item["type"]
        if kind.startswith("h"):
            doc.add_heading(item["text"], level=int(kind[1:]))
        elif kind == "p":
            para = doc.add_paragraph(item["text"])
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            para.paragraph_format.widow_control = True
        elif kind == "bullets":
            for entry in item["items"]:
                para = doc.add_paragraph(style="List Bullet")
                para.add_run(entry)
        elif kind == "figure":
            add_figure_zh(doc, item["number"])
        elif kind == "table":
            add_table_zh(doc, item["number"])

    doc.add_heading("参考文献", level=1)
    for ref in base.REFERENCES:
        para = doc.add_paragraph(ref)
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.left_indent = Inches(0.25)
        para.paragraph_format.first_line_indent = Inches(-0.25)
        para.paragraph_format.space_after = Pt(5)
        para.paragraph_format.line_spacing = 1.1
    doc.save(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    base._ORIGINAL_SET_RUN_FONT = base.set_run_font
    base._ORIGINAL_CONFIGURE_STYLES = base.configure_styles
    base._ORIGINAL_CONFIGURE_PAGE = base.configure_page
    base._ORIGINAL_LOAD_TABLE = base.load_table
    base.set_run_font = set_run_font_zh

    prepare_muse_idr_figures()
    MARKDOWN_PATH.write_text(render_markdown(), encoding="utf-8", newline="\n")
    build_docx(DOCX_PATH)
    report = {
        "schema_version": 1,
        "status": "pass",
        "experiment": "c2_chinese_manuscript",
        "model_name": "MUSE-IDR",
        "internal_locked_identifier": "A10-locked",
        "source_english_docx": str(base.OUT / "A10_locked_manuscript.docx"),
        "source_english_docx_sha256": sha256_file(base.OUT / "A10_locked_manuscript.docx"),
        "docx": str(DOCX_PATH.relative_to(ROOT)).replace("\\", "/"),
        "docx_sha256": sha256_file(DOCX_PATH),
        "markdown": str(MARKDOWN_PATH.relative_to(ROOT)).replace("\\", "/"),
        "markdown_sha256": sha256_file(MARKDOWN_PATH),
        "main_figures": 3,
        "main_figure_sha256": {
            str(number): sha256_file(path) for number, path in FIGURE_FILES.items()
        },
        "main_tables": 2,
        "references": len(base.REFERENCES),
        "author_placeholders_remaining": True,
        "caid_labels_used_for_training_or_tuning": False,
    }
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8", newline="\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
