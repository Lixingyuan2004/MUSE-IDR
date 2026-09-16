from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "outputs" / "paper" / "c2_full_manuscript"
C1 = ROOT / "outputs" / "paper" / "c1_manuscript_materials"
RELEASE = ROOT / "release"

TITLE = (
    "A leakage-audited cross-architecture protein language model ensemble "
    "for intrinsic disorder prediction"
)
RUNNING_TITLE = "Locked ESM-2 ensemble for IDR prediction"

# Documents skill design choice: narrative_proposal preset with one named,
# consistently reused override, journal_title_page_compact, derived from the
# editorial_cover pattern. No generic Word defaults are relied upon.
DESIGN = {
    "preset": "narrative_proposal",
    "header_pattern": "editorial_cover",
    "named_override": "journal_title_page_compact",
    "page": {
        "width_in": 8.5,
        "height_in": 11.0,
        "margins_in": 1.0,
        "header_footer_in": 0.492,
        "content_width_dxa": 9360,
    },
    "body": {
        "font": "Calibri",
        "size_pt": 11,
        "alignment": "justified",
        "before_pt": 0,
        "after_pt": 8,
        "line_spacing": 1.333,
    },
    "headings": {
        "h1": {"size_pt": 16, "color": "2E74B5", "before_pt": 18, "after_pt": 10},
        "h2": {"size_pt": 13, "color": "2E74B5", "before_pt": 12, "after_pt": 6},
        "h3": {"size_pt": 12, "color": "1F4D78", "before_pt": 8, "after_pt": 4},
    },
    "lists": {
        "marker_aligned_in": 0.181,
        "text_indent_in": 0.375,
        "hanging_in": 0.194,
        "after_pt": 4,
        "line_spacing": 1.208,
    },
    "tables": {
        "width_dxa": 9360,
        "indent_dxa": 120,
        "cell_margins_dxa": {"top": 80, "bottom": 80, "start": 120, "end": 120},
        "header_fill": "F4F6F9",
    },
    "journal_title_page_compact": {
        "title_size_pt": 20,
        "title_color": "203748",
        "kicker_color": "7A5A00",
        "subtitle_color": "4B6575",
        "top_spacer_pt": 38,
        "no_decorative_rule": True,
    },
}

ABSTRACT = (
    "Intrinsic-disorder predictors are commonly compared on benchmarks whose operational "
    "definitions and residue prevalence differ, making broad superiority claims fragile. "
    "We developed A10-locked, a leakage-audited ensemble that combines two frozen ESM-2-650M "
    "heads: a final-layer multiscale convolutional model and an otherwise matched model with "
    "a learned mixture of the final four representation layers. Thirty members spanning two "
    "architectures, five homology-isolated folds and three seeds are averaged uniformly in "
    "logit space. The model was finalized before CAID2 or CAID3 label access. On the PDB tracks, "
    "A10 achieved ROC-AUC/AUPRC values of 0.958/0.934 (CAID2) and 0.954/0.931 (CAID3), exceeding "
    "a reproduced LoRA-DR-Suite 650M baseline by 0.037 and 0.033 ROC-AUC, respectively, in paired "
    "whole-protein bootstrap analyses. On NOX, performance was lower (0.781/0.361 and 0.852/0.606), "
    "and A10 was not uniformly superior. It remained statistically comparable to reproduced "
    "PUNCH2-Light Released-13 across all four tracks. Post-lock diagnostics associated A10 with "
    "low error on long disordered segments but excess false positives in low-disorder and long "
    "ordered contexts. These results support a track-specific view of predictor quality and show "
    "how model locking, homology control and paired protein-level uncertainty can make external "
    "benchmarking more interpretable."
)

KEYWORDS = [
    "intrinsic disorder",
    "protein language models",
    "ESM-2",
    "ensemble learning",
    "CAID",
]


def h(level: int, text: str) -> dict:
    return {"type": f"h{level}", "text": text}


def p(text: str) -> dict:
    return {"type": "p", "text": text}


def bullets(items: list[str]) -> dict:
    return {"type": "bullets", "items": items}


def figure(number: int) -> dict:
    return {"type": "figure", "number": number}


def table(number: int) -> dict:
    return {"type": "table", "number": number}


CONTENT = [
    h(1, "1 Introduction"),
    p(
        "Intrinsically disordered proteins and regions do not occupy a single stable tertiary "
        "structure under physiological conditions, yet they participate in signalling, regulation, "
        "molecular recognition and assembly (Wright and Dyson, 2015). Their conformational "
        "heterogeneity and context dependence make experimental characterization difficult and "
        "increase reliance on sequence-based prediction. DisProt provides manually curated, "
        "experimentally supported annotations, while the Critical Assessment of protein Intrinsic "
        "Disorder prediction (CAID) evaluates methods on previously unseen proteins (Piovesan et "
        "al., 2017; Necci et al., 2021)."
    ),
    p(
        "Recent CAID rounds have shown both rapid progress and persistent benchmark dependence. "
        "The Disorder-NOX reference treats residues without qualifying disorder evidence as ordered, "
        "whereas Disorder-PDB masks negative residues lacking observed structural coordinates. The "
        "resulting tracks differ in prevalence and in what constitutes reliable negative evidence "
        "(Del Conte et al., 2023; Mehdiabadi et al., 2026). A method can therefore rank strongly on "
        "one track and less strongly on the other without either result being erroneous. Reporting a "
        "single pooled score risks obscuring this operational heterogeneity."
    ),
    p(
        "Protein language models (PLMs) provide contextual residue representations learned from "
        "large sequence collections. ESM-2 representations encode evolutionary and structural "
        "regularities and have become a common input to disorder predictors (Lin et al., 2023; "
        "Mehdiabadi et al., 2026). Existing methods exploit them in different ways. LoRA-DR-Suite "
        "adapts attention projections of PLMs with low-rank updates, while PUNCH2-Light combines "
        "ProtT5 and one-hot features through deep convolutional ensembles (Lombardi et al., 2025; "
        "Meng and Pollastri, 2025). These approaches motivate a complementary question: how much "
        "can be obtained from a fully frozen backbone when layer and local-context diversity are "
        "combined under a strict leakage-control protocol?"
    ),
    p(
        "We address this question with A10-locked, a 30-member ensemble of two heads over frozen "
        "ESM-2-650M residue representations. One family uses the final representation layer and a "
        "parallel multiscale convolutional head; the second learns a scalar mixture of layers 30-33 "
        "before the same contextual stage. All folds, seeds, architectures and aggregation weights "
        "were fixed before CAID2 or CAID3 labels were accessed. Our aims were to (i) quantify external "
        "performance on both CAID2 and CAID3 Disorder-NOX and Disorder-PDB tracks, (ii) compare A10 "
        "with locally reproduced, residue-aligned LoRA-DR-Suite and PUNCH2-Light baselines, and "
        "(iii) identify the protein and segment regimes that explain track-specific performance."
    ),

    h(1, "2 Methods"),
    h(2, "2.1 Study design and model locking"),
    p(
        "The study was organized as a development phase followed by a locked external evaluation. "
        "Model architecture, training data, cross-validation partitions, random seeds, ensemble "
        "membership, member weights, probability threshold and output definition were finalized "
        "before external labels were read. The locked artifact contains 30 checkpoint hashes: 15 A2 "
        "members and 15 A8 members from five folds and seeds 17, 29 and 43. Each member contributes "
        "one logit with weight 1/30. No out-of-fold or CAID labels were used to fit ensemble weights, "
        "recalibrate probabilities or select a track-specific variant."
    ),
    h(2, "2.2 Training data and leakage controls"),
    p(
        "The training corpus contains 1,133 proteins and 573,421 residues. Of these, 202,662 residues "
        "have known binary labels (105,036 disordered and 97,626 ordered); all other residues are "
        "marked unknown and excluded from loss and metrics. Positive labels were derived from direct "
        "classic-disorder evidence in the historical DisProt 2018_11 snapshot. Negative labels "
        "required experimentally observed Protein Data Bank coordinates mapped back to the candidate "
        "sequence. X-ray structures were restricted to resolution <=3.0 A and cryo-electron "
        "microscopy structures to <=4.0 A; nuclear magnetic resonance structures had no resolution "
        "cutoff. Regions annotated as disorder, molten globule or pre-molten globule, together with a "
        "five-residue boundary margin, could not become negatives."
    ),
    p(
        "CAID2/CAID3 exclusion was applied before cross-validation at three levels: matching DisProt "
        "or UniProt identifiers and isoform roots, exact normalized-sequence SHA256, and homology. "
        "The homology criterion removed a candidate when MMseqs2 reported >30% identity and at least "
        "80% coverage of both query and target. MMseqs2 release 18-8cc5c was used and all raw search "
        "outputs and hashes were retained (Steinegger and Soding, 2017). After exclusion, an internal "
        "all-versus-all graph was built with the same criterion. Its 984 connected components (190 "
        "non-self edges) were assigned intact to five folds while balancing positive residues, "
        "negative residues, protein count and total length. The fold assignment used seed 17 and was "
        "reused for all model-initialization seeds; no qualifying homology edge crossed folds."
    ),
    h(2, "2.3 Frozen ESM-2 representations"),
    p(
        "All models used facebook/esm2_t33_650M_UR50D at revision "
        "08e4846e537177426273712802403f7ba8261b6c. Backbone parameters were frozen. Sequences longer "
        "than the 1,022-residue content limit were processed with overlapping windows at stride 511; "
        "overlap logits or representations were fused with a deterministic centre-weighted average. "
        "A2 cached the final-layer representation [L,1280]. A8 independently generated and hash-"
        "verified float16 caches for layers 30, 31, 32 and 33 with shape [L,4,1280]. Caches contained "
        "no residue labels."
    ),
    h(2, "2.4 A2 and A8 prediction heads"),
    p(
        "A2 maps each 1,280-dimensional residue vector through LayerNorm and a 256-dimensional GELU "
        "projection. Four depthwise-separable one-dimensional convolution branches use kernels 3, 7, "
        "15 and 31. Their outputs are fused through a learned gate and residual connection, followed "
        "by dropout (0.2) and a single residue logit. This explicitly supplies multiple local context "
        "scales without updating ESM-2."
    ),
    p(
        "A8 has the same contextual head but first computes a global learned softmax mixture of ESM-2 "
        "layers 30-33. The mixture was initialized to [0.1, 0.1, 0.1, 0.7], maintaining an initial "
        "bias toward the final layer. It adds only four trainable scalar parameters relative to A2. "
        "Both families have one output head representing classic intrinsic-disorder probability; "
        "neither predicts soft disorder or track-specific NOX/PDB outputs."
    ),
    figure(1),
    h(2, "2.5 Training and ensemble construction"),
    p(
        "Heads were trained for at most 20 epochs with AdamW, learning rate 3x10^-4, weight decay "
        "0.01, protein batch size 4 and gradient-norm clipping at 1.0. The A8 layer mixture used a "
        "separate learning rate of 3x10^-3. Binary cross-entropy with logits was evaluated only on "
        "known residues and used a fold-specific positive-class weight equal to the negative-to-"
        "positive residue ratio. Validation micro ROC-AUC controlled checkpointing and early stopping "
        "with patience four. For external inference, the 30 member logits were averaged uniformly and "
        "then transformed with the sigmoid function."
    ),
    h(2, "2.6 External datasets and reproduced baselines"),
    p(
        "External evaluation used CAID2 and CAID3 Disorder-NOX and Disorder-PDB references. CAID2-"
        "NOX contained 210 proteins and 160,802 known residues; CAID2-PDB contained 348 proteins and "
        "130,877 known residues. CAID3-NOX contained 204 proteins and 99,977 known residues; CAID3-"
        "PDB contained 319 proteins and 99,239 known residues. Masked labels were excluded. A10 "
        "predictions were generated once from sequence-only FASTA files, hash-locked and then joined "
        "to each reference."
    ),
    p(
        "We reproduced three eligible baseline variants. LoRA-DR-Suite used the authors' ESM-2-650M "
        "adapter trained on DisProt 7 only, excluding variants trained with CAID labels and all "
        "soft-disorder heads. The released adapter, saved token-classification head and query/value "
        "LoRA tensors were verified before inference. PUNCH2-Light was reproduced from commit "
        "6c7935b3597c056d2e6b3845bb54fc101c5bc574 using official one-hot and ProtT5 features. Because "
        "the paper specifies eight ensemble members whereas the released entry point loads 13, both "
        "Paper-8 and Released-13 variants were reported separately. All comparator predictions were "
        "aligned residue-by-residue to the same references and hash-locked before evaluation."
    ),
    h(2, "2.7 Metrics and confirmatory statistics"),
    p(
        "We report full-precision ROC-AUC, trapezoidal area under the precision-recall curve (AUPRC), "
        "average precision score (APS), F1 and Matthews correlation coefficient at the predeclared "
        "threshold 0.5, Fmax, and Brier score. Precision-recall metrics were retained because class "
        "imbalance differs between tracks (Davis and Goadrich, 2006; Saito and Rehmsmeier, 2015). "
        "The complete protein was the resampling unit, preserving within-protein residue dependence."
    ),
    p(
        "For each track, 2,000 paired bootstrap resamples were shared by all four models. Percentile "
        "95% confidence intervals and two-sided bootstrap p-values were computed for A10-minus-"
        "comparator metric differences. Paired DeLong tests were additionally calculated for ROC-AUC "
        "(DeLong et al., 1988). Holm correction controlled 48 primary hypotheses: A10 versus LoRA-DR-"
        "Suite and A10 versus PUNCH2-Light Released-13 across four tracks and six metrics. Paper-8 "
        "comparisons formed a separate 24-test supplementary family (Holm, 1979)."
    ),
    h(2, "2.8 Post-lock descriptive analysis"),
    p(
        "Error and complementarity analyses were descriptive and did not introduce new confirmatory "
        "tests. Predeclared protein strata were length (1-200, 201-500, 501-1000 and 1001+ residues) "
        "and known-label disorder fraction (0-0.10, (0.10,0.30], (0.30,0.60] and (0.60,1.00]). "
        "Residue diagnostics used distance to the nearest observed 0/1 transition and true ordered or "
        "disordered run lengths (1-15, 16-30, 31-100 and 101+ residues). Threshold-error disagreement, "
        "model-only correctness and Spearman score correlation characterized complementarity. None of "
        "these results was used to alter A10."
    ),
    h(2, "2.9 Reproducibility and software"),
    p(
        "All source archives, model checkpoints, predictions, external references and analysis outputs "
        "were protected by SHA256 manifests. Locked-inference smoke tests required all 30 members, "
        "excluded the development checkpoint, verified one probability per input residue and produced "
        "identical predictions on cache reuse. Internal CAID-compatible metrics exactly matched the "
        "official CAID implementation at revision 7bab7e8880d7f949e480fdb377c5a5c8546ae336 for "
        "both CAID3 tracks. Analyses used Python 3.12, PyTorch 2.12.1, Transformers 5.15.1 and MMseqs2 "
        "18-8cc5c. Training and inference ran on an NVIDIA GeForce RTX 4090."
    ),

    h(1, "3 Results"),
    h(2, "3.1 Locked ensemble and external performance"),
    p(
        "A10 combined two representation views rather than fitting a larger track-specific head. "
        "Fifteen A2 members used the final ESM-2 layer, and 15 A8 members learned a mixture of the "
        "final four layers. The fixed equal-logit ensemble achieved ROC-AUC/AUPRC values of "
        "0.958/0.934 on CAID2-PDB and 0.954/0.931 on CAID3-PDB. Performance on NOX was lower and "
        "dataset dependent: 0.781/0.361 on CAID2 and 0.852/0.606 on CAID3 (Table 1; Fig. 2)."
    ),
    table(1),
    figure(2),
    h(2, "3.2 Confirmatory paired comparisons"),
    p(
        "Against LoRA-DR-Suite 650M, A10 had higher ROC-AUC on both PDB tracks. The paired differences "
        "were +0.0374 on CAID2-PDB (95% protein-bootstrap CI +0.0227 to +0.0534; Holm-adjusted "
        "bootstrap p=0.048) and +0.0334 on CAID3-PDB (+0.0192 to +0.0484; adjusted p=0.048). On "
        "CAID2-NOX, A10 was lower by 0.0541 (-0.0973 to -0.0160), although the 48-test Holm-adjusted "
        "bootstrap p-value was 0.180. The CAID3-NOX difference was -0.0079 and crossed zero."
    ),
    p(
        "A10 and PUNCH2-Light Released-13 had overlapping paired protein-bootstrap intervals on all "
        "four tracks. A10 point estimates were higher by 0.0013 and 0.0078 on CAID2-NOX and PDB and by "
        "0.0122 and 0.0011 on CAID3-NOX and PDB, respectively. These results support comparable "
        "external performance rather than universal superiority. Supplementary Paper-8 comparisons "
        "showed similar directions but were interpreted within their separate correction family "
        "(Table 2; Fig. 3)."
    ),
    table(2),
    figure(3),
    h(2, "3.3 Disorder-fraction and segment-length regimes"),
    p(
        "Post-lock stratification localized the largest weaknesses to proteins with 0-10% known "
        "disordered residues. In this stratum, A10-minus-LoRA ROC-AUC differences were approximately "
        "-0.090 on CAID2-NOX and -0.058 on CAID3-NOX. Direction reversed in higher-disorder strata, "
        "including +0.105 in the CAID2-NOX 30-60% stratum and +0.085 in the CAID3-NOX >60% stratum. "
        "These exploratory values are visualized in Supplementary Fig. S1 and tabulated in "
        "Supplementary Table S2."
    ),
    p(
        "For true disordered runs of at least 101 residues, A10's fixed-threshold error rate ranged "
        "from approximately 0.047 to 0.072 across the four tracks. The corresponding ranges were "
        "0.226-0.311 for LoRA-DR-Suite and 0.120-0.133 for PUNCH2-Light Released-13. Conversely, A10 "
        "made more false-positive errors in long ordered segments on NOX. On CAID2-NOX, its mean "
        "predicted probability was 0.515 despite prevalence 0.195, consistent with a positive "
        "calibration bias (Supplementary Fig. S2; Supplementary Tables S3-S4)."
    ),
    h(2, "3.4 Residue-level complementarity"),
    p(
        "A10 and the reproduced baselines did not make identical errors. Relative to LoRA-DR-Suite, "
        "A10 alone was correct for 8.8% and 7.8% of residues on CAID2-PDB and CAID3-PDB, while the "
        "comparator alone was correct for 2.6% and 4.3%. The direction reversed on NOX: LoRA-only "
        "correctness exceeded A10-only correctness. A10 and PUNCH2-Light scores were highly correlated "
        "(Spearman approximately 0.83-0.90), yet their thresholded predictions still disagreed on "
        "5.6-11.8% of residues depending on track and variant (Supplementary Table S5)."
    ),

    h(1, "4 Discussion"),
    h(2, "4.1 Principal findings"),
    p(
        "A10-locked is best understood as a leakage-audited, track-aware external validation of a "
        "cross-architecture frozen-PLM ensemble. It is particularly effective on Disorder-PDB and "
        "long disordered segments, and its PDB advantage over the reproduced LoRA-DR-Suite baseline "
        "remained after paired whole-protein resampling and family-wise correction. At the same time, "
        "it did not dominate Disorder-NOX and was statistically comparable to PUNCH2-Light Released-"
        "13 across the four principal track/dataset combinations. This bounded claim is stronger and "
        "more reproducible than a universal state-of-the-art statement."
    ),
    h(2, "4.2 Why track definitions matter"),
    p(
        "The NOX/PDB contrast is central to interpretation. Disorder-NOX penalizes high scores outside "
        "annotated disorder even when structural negative evidence is incomplete; Disorder-PDB masks "
        "many such residues and focuses evaluation on experimentally supported positives and observed "
        "ordered negatives. A10's tendency to assign sustained high scores to locally disorder-like "
        "ordered regions therefore harms low-prevalence NOX proteins but helps retain long IDRs. The "
        "same behavior can be advantageous or detrimental under different operational references. "
        "CAID's multi-track design is consequently informative rather than redundant."
    ),
    h(2, "4.3 Architectural interpretation"),
    p(
        "The ensemble combines two economical forms of diversity. A2 explicitly integrates local "
        "contexts at four scales from the final ESM-2 layer. A8 exposes the same head to a learned "
        "mixture of the last four layers, adding only four trainable scalars. Averaging across "
        "architectures, homology-isolated folds and initializations reduces dependence on a single "
        "representation or data partition. The low error on long disordered runs is consistent with "
        "multiscale context aggregation, but external results alone do not establish causality; a "
        "component-level ablation on a newly locked development protocol would be needed."
    ),
    h(2, "4.4 Relation to reproduced baselines"),
    p(
        "LoRA-DR-Suite adapts the PLM backbone and is stronger on NOX, suggesting that task-specific "
        "attention updates can improve discrimination or calibration when disorder prevalence is low. "
        "A10 instead keeps the backbone frozen and is stronger on the two PDB tracks. PUNCH2-Light "
        "uses ProtT5 plus one-hot inputs and is close to A10 in aggregate, yet the residual disagreement "
        "indicates that representation families remain complementary. Future label-free stacking, "
        "gating or distillation is plausible, but any changed predictor must receive a new identifier "
        "and be evaluated on untouched external data rather than presented as a revision of A10."
    ),
    h(2, "4.5 Limitations"),
    bullets([
        "The external evidence is restricted to CAID2 and CAID3 classic-disorder tracks and does not establish performance for binding disorder, linkers or soft disorder.",
        "NOX and PDB are different operational references; neither can be treated as a universally complete biological ground truth.",
        "The 30-member ensemble shares backbone extraction but remains more costly to store and evaluate than a single prediction head.",
        "Fixed-threshold and Brier analyses reveal overprediction on low-prevalence NOX data; external-label recalibration was intentionally prohibited.",
        "Confirmatory comparisons cover locally reproduced LoRA-DR-Suite and PUNCH2-Light variants. Contextual official rankings are not equivalent to residue-aligned local reproduction.",
        "Subgroup, boundary, segment and complementarity analyses are descriptive and should not be interpreted as new confirmatory hypotheses.",
    ]),
    h(2, "4.6 Future work"),
    p(
        "A successor should target the identified failure regime using training-set cross-validation "
        "only. Candidate interventions include stronger sampling of experimentally ordered negatives, "
        "calibration regularization, boundary-aware objectives, or a lightweight gate that suppresses "
        "global overprediction in low-disorder proteins. Distillation could reduce the 30-head cost. "
        "The current complementarity with LoRA-DR-Suite and PUNCH2-Light also motivates a predeclared "
        "cross-family ensemble, provided its weights are learned without CAID2/CAID3 labels and its "
        "external test set is genuinely untouched."
    ),

    h(1, "5 Conclusion"),
    p(
        "A10-locked combines multiscale final-layer and learned multilayer ESM-2 heads under a strict "
        "homology, hashing and model-locking protocol. It performs strongly on CAID2/CAID3 Disorder-"
        "PDB and long disordered segments, remains comparable to reproduced PUNCH2-Light, and exposes "
        "a clear weakness on low-disorder NOX proteins. The main contribution is therefore both a "
        "model and an evaluation discipline: external labels are reserved for post-lock assessment, "
        "uncertainty is paired at the protein level, and conclusions follow benchmark definitions "
        "rather than collapsing them into a single claim."
    ),

    h(1, "Data availability"),
    p(
        "CAID reference data and prediction-format specifications are available from the CAID project. "
        "DisProt and PDB source data are available from their respective public resources. The exact "
        "historical snapshots, derived training manifests, model checkpoints and result archives used "
        "in this study are identified by SHA256 manifests in the project artifact. [REPOSITORY DOI/URL "
        "TO BE INSERTED BEFORE SUBMISSION]. Third-party weights are not redistributed where license "
        "terms are absent or unclear."
    ),
    h(1, "Code availability"),
    p(
        "Training, locked inference, baseline reproduction, metric calculation, paired bootstrap and "
        "manuscript-material generation scripts will be released at [CODE REPOSITORY URL/DOI]. The "
        "release should include the environment specification, source-revision locks and a minimal "
        "sequence-only inference example."
    ),
    h(1, "Acknowledgements"),
    p("[ACKNOWLEDGEMENTS TO BE COMPLETED BY THE AUTHORS]."),
    h(1, "Funding"),
    p("[FUNDING SOURCES AND GRANT NUMBERS TO BE COMPLETED BY THE AUTHORS]."),
    h(1, "Conflict of interest"),
    p("The authors declare [NO CONFLICTS / CONFLICTS TO BE COMPLETED BY THE AUTHORS]."),
    h(1, "Author contributions"),
    p("[CRediT CONTRIBUTOR ROLES TO BE COMPLETED AFTER THE AUTHOR LIST IS FINALIZED]."),
    h(1, "AI-assisted drafting disclosure"),
    p(
        "OpenAI Codex was used to assist source organization, language drafting and programmatic Word "
        "assembly. It did not select models, access evaluation labels during model development, or "
        "interpret unreviewed raw experimental data. All scientific statements, citations, analyses "
        "and final wording require verification and approval by the human authors before submission."
    ),
]


REFERENCES = [
    "Davis,J. and Goadrich,M. (2006) The relationship between Precision-Recall and ROC curves. In: Proceedings of the 23rd International Conference on Machine Learning, pp. 233-240. doi:10.1145/1143844.1143874.",
    "Del Conte,A., Mehdiabadi,M., Bouhraoua,A., Monzon,A.M., Tosatto,S.C.E. and Piovesan,D. (2023) Critical assessment of protein intrinsic disorder prediction (CAID) - Results of round 2. Proteins, 91, 1925-1934. doi:10.1002/prot.26582.",
    "DeLong,E.R., DeLong,D.M. and Clarke-Pearson,D.L. (1988) Comparing the areas under two or more correlated receiver operating characteristic curves: a nonparametric approach. Biometrics, 44, 837-845. doi:10.2307/2531595.",
    "Elnaggar,A. et al. (2022) ProtTrans: Toward understanding the language of life through self-supervised learning. IEEE Transactions on Pattern Analysis and Machine Intelligence, 44, 7112-7127. doi:10.1109/TPAMI.2021.3095381.",
    "Hatos,A. et al. (2020) DisProt: intrinsic protein disorder annotation in 2020. Nucleic Acids Research, 48, D269-D276. doi:10.1093/nar/gkz975.",
    "Holm,S. (1979) A simple sequentially rejective multiple test procedure. Scandinavian Journal of Statistics, 6, 65-70.",
    "Hu,E.J. et al. (2022) LoRA: Low-rank adaptation of large language models. In: International Conference on Learning Representations.",
    "Lin,Z. et al. (2023) Evolutionary-scale prediction of atomic-level protein structure with a language model. Science, 379, 1123-1130. doi:10.1126/science.ade2574.",
    "Lombardi,G., Seoane,B. and Carbone,A. (2025) LoRA-DR-suite: adapted embeddings predict intrinsic and soft disorder from protein sequences. Bioinformatics, 41, i439-i448. doi:10.1093/bioinformatics/btaf185.",
    "Mehdiabadi,M., Del Conte,A., Nugnes,M.V., Aspromonte,M.C., Tosatto,S.C.E. and Piovesan,D. (2026) Critical Assessment of Protein Intrinsic Disorder Round 3 - Predicting Disorder in the Era of Protein Language Models. Proteins, 94, 414-424. doi:10.1002/prot.70045.",
    "Meng,D. and Pollastri,G. (2025) PUNCH2: Explore the strategy for intrinsically disordered protein predictor. PLOS ONE, 20, e0319208. doi:10.1371/journal.pone.0319208.",
    "Necci,M., Piovesan,D., CAID Predictors, DisProt Curators and Tosatto,S.C.E. (2021) Critical assessment of protein intrinsic disorder prediction. Nature Methods, 18, 472-481. doi:10.1038/s41592-021-01117-3.",
    "Piovesan,D. et al. (2017) DisProt 7.0: a major update of the database of disordered proteins. Nucleic Acids Research, 45, D219-D227. doi:10.1093/nar/gkw1056.",
    "Saito,T. and Rehmsmeier,M. (2015) The precision-recall plot is more informative than the ROC plot when evaluating binary classifiers on imbalanced datasets. PLOS ONE, 10, e0118432. doi:10.1371/journal.pone.0118432.",
    "Steinegger,M. and Soding,J. (2017) MMseqs2 enables sensitive protein sequence searching for the analysis of massive data sets. Nature Biotechnology, 35, 1026-1028. doi:10.1038/nbt.3988.",
    "Wright,P.E. and Dyson,H.J. (2015) Intrinsically disordered proteins in cellular signalling and regulation. Nature Reviews Molecular Cell Biology, 16, 18-29. doi:10.1038/nrm3920.",
]

FIGURE_LEGENDS = {
    1: (
        "Locked A10 architecture and evaluation protocol. Frozen ESM-2-650M residue "
        "representations feed two model families. A2 uses the final representation layer with a "
        "multiscale context head; A8 learns a mixture of layers 30-33 before contextual prediction. "
        "Three seeds and five folds per family yield 30 logits, combined by a fixed uniform logit "
        "mean. All model choices were fixed before CAID2/CAID3 label access."
    ),
    2: (
        "External-test performance on CAID2 and CAID3. Full-precision ROC-AUC and AUPRC for "
        "A10-locked and three reproduced baseline variants. All models have complete residue "
        "coverage in the aligned comparison."
    ),
    3: (
        "Paired protein-bootstrap differences in ROC-AUC. Points show A10 minus comparator; "
        "horizontal lines show percentile 95% confidence intervals from 2,000 paired whole-protein "
        "bootstrap replicates. Asterisks denote Holm-adjusted bootstrap p<0.05. Paper-8 comparisons "
        "belong to a separate supplementary correction family."
    ),
}

FIGURE_FILES = {
    1: C1 / "figure_1_locked_a10_architecture.png",
    2: C1 / "figure_2_external_performance.png",
    3: C1 / "figure_3_paired_roc_forest.png",
}


def sha256_file(path: Path) -> str:
    hsh = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            hsh.update(chunk)
    return hsh.hexdigest()


def word_count(text: str) -> int:
    return len(re.findall(r"\b[\w'-]+\b", text))


def load_table(number: int) -> list[list[str]]:
    path = C1 / (
        "table_1_external_performance.csv"
        if number == 1
        else "table_2_paired_roc_statistics.csv"
    )
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        raw = [row for row in csv.reader(handle)]
    if number == 1:
        selected = [[row[i] for i in (0, 1, 2, 6, 7, 8, 9, 10)] for row in raw]
        selected[0] = [
            "Dataset", "Track", "Model", "ROC-AUC", "AUPRC", "APS",
            "F1@0.5", "MCC@0.5",
        ]
        for row in selected[1:]:
            for idx in range(3, 8):
                row[idx] = f"{float(row[idx]):.3f}"
        return selected
    selected = [[row[i] for i in (0, 1, 2, 4, 5, 6, 7, 8)] for row in raw]
    selected[0] = [
        "Dataset", "Track", "Comparator", "Proteins", "A10 delta",
        "95% CI", "Bootstrap Holm p", "DeLong Holm p",
    ]
    for row in selected[1:]:
        row[3] = str(int(float(row[3])))
        row[4] = f"{float(row[4]):+.4f}"
        bounds = json.loads(row[5])
        row[5] = f"[{bounds[0]:+.4f}, {bounds[1]:+.4f}]"
        row[6] = f"{float(row[6]):.3g}"
        row[7] = f"{float(row[7]):.3g}"
    return selected


def set_run_font(run, name: str = "Calibri", size: float | None = None,
                 color: str | None = None, bold: bool | None = None,
                 italic: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_row_cant_split(row) -> None:
    """Keep every logical table row on one page in Word/PDF rendering."""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def apply_table_geometry(table_obj, widths: list[int]) -> None:
    assert sum(widths) == DESIGN["tables"]["width_dxa"]
    table_obj.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_obj.autofit = False
    tbl_pr = table_obj._tbl.tblPr
    for tag in ("w:tblW", "w:tblInd"):
        existing = tbl_pr.find(qn(tag))
        if existing is not None:
            tbl_pr.remove(existing)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(DESIGN["tables"]["width_dxa"]))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(DESIGN["tables"]["indent_dxa"]))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)

    grid = table_obj._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    margins = DESIGN["tables"]["cell_margins_dxa"]
    for row in table_obj.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[idx] / 1440)
            set_cell_margins(cell, **margins)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color="6B7280")
    for kind, text in (
        ("begin", None),
        ("instr", " PAGE "),
        ("separate", None),
        ("text", "1"),
        ("end", None),
    ):
        field_run = OxmlElement("w:r")
        if kind in {"begin", "separate", "end"}:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        elif kind == "instr":
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = text
        else:
            node = OxmlElement("w:t")
            node.text = text
        field_run.append(node)
        paragraph._p.append(field_run)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = DESIGN["body"]["font"]
    normal._element.rPr.rFonts.set(qn("w:ascii"), DESIGN["body"]["font"])
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), DESIGN["body"]["font"])
    normal.font.size = Pt(DESIGN["body"]["size_pt"])
    pf = normal.paragraph_format
    pf.space_before = Pt(DESIGN["body"]["before_pt"])
    pf.space_after = Pt(DESIGN["body"]["after_pt"])
    pf.line_spacing = DESIGN["body"]["line_spacing"]
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for level in (1, 2, 3):
        style = doc.styles[f"Heading {level}"]
        token = DESIGN["headings"][f"h{level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(token["size_pt"])
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(token["color"])
        style.paragraph_format.space_before = Pt(token["before_pt"])
        style.paragraph_format.space_after = Pt(token["after_pt"])
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string("374151")
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.line_spacing = 1.05

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(DESIGN["lists"]["text_indent_in"])
        style.paragraph_format.first_line_indent = Inches(-DESIGN["lists"]["hanging_in"])
        style.paragraph_format.space_after = Pt(DESIGN["lists"]["after_pt"])
        style.paragraph_format.line_spacing = DESIGN["lists"]["line_spacing"]


def configure_page(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(DESIGN["page"]["width_in"])
        section.page_height = Inches(DESIGN["page"]["height_in"])
        section.top_margin = Inches(DESIGN["page"]["margins_in"])
        section.right_margin = Inches(DESIGN["page"]["margins_in"])
        section.bottom_margin = Inches(DESIGN["page"]["margins_in"])
        section.left_margin = Inches(DESIGN["page"]["margins_in"])
        section.header_distance = Inches(DESIGN["page"]["header_footer_in"])
        section.footer_distance = Inches(DESIGN["page"]["header_footer_in"])
        section.different_first_page_header_footer = True

        header = section.header
        hp = header.paragraphs[0]
        hp.clear()
        hp.paragraph_format.space_after = Pt(0)
        hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
        left = hp.add_run(RUNNING_TITLE)
        set_run_font(left, size=8.5, color="6B7280")
        right = hp.add_run("\tOriginal Article | Draft")
        set_run_font(right, size=8.5, color="6B7280")

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.clear()
        add_page_field(fp)


def add_title_page(doc: Document) -> None:
    override = DESIGN["journal_title_page_compact"]
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(override["top_spacer_pt"])

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(12)
    r = kicker.add_run("ORIGINAL ARTICLE | AUTHOR DRAFT")
    set_run_font(r, size=9.5, color=override["kicker_color"], bold=True)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(10)
    title_p.paragraph_format.keep_with_next = True
    r = title_p.add_run(TITLE)
    set_run_font(
        r,
        size=override["title_size_pt"],
        color=override["title_color"],
        bold=True,
    )

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(22)
    r = subtitle.add_run(
        "Locked external validation on CAID2 and CAID3 with paired protein-level uncertainty"
    )
    set_run_font(r, size=12.5, color=override["subtitle_color"], italic=True)

    for text, size, bold in (
        ("[AUTHOR NAMES AND ORCID IDENTIFIERS]", 11.5, True),
        ("[AFFILIATIONS]", 10.5, False),
        ("[CORRESPONDING AUTHOR ADDRESS AND EMAIL]", 10.0, False),
    ):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(5)
        run = para.add_run(text)
        set_run_font(run, size=size, color="374151", bold=bold)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(22)
    meta.paragraph_format.space_after = Pt(4)
    run = meta.add_run("Target journal: Bioinformatics Advances")
    set_run_font(run, size=10, color="203748", bold=True)

    meta2 = doc.add_paragraph()
    meta2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta2.add_run(f"Running head: {RUNNING_TITLE} | Draft date: {date.today().isoformat()}")
    set_run_font(run, size=9.5, color="6B7280")
    doc.add_page_break()


def add_abstract(doc: Document) -> None:
    doc.add_heading("Abstract", level=1)
    para = doc.add_paragraph(ABSTRACT)
    para.paragraph_format.keep_together = True
    key = doc.add_paragraph()
    key.paragraph_format.space_before = Pt(3)
    label = key.add_run("Keywords: ")
    set_run_font(label, bold=True)
    value = key.add_run("; ".join(KEYWORDS))
    set_run_font(value)


def add_body_paragraph(doc: Document, text: str) -> None:
    para = doc.add_paragraph(text)
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para.paragraph_format.widow_control = True


def add_bullet_list(doc: Document, items: list[str]) -> None:
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        para.add_run(item)


def add_figure(doc: Document, number: int) -> None:
    image_path = FIGURE_FILES[number]
    if not image_path.is_file():
        raise FileNotFoundError(image_path)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(8)
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.keep_with_next = True
    run = para.add_run()
    run.add_picture(str(image_path), width=Inches(6.15))
    shape = doc.inline_shapes[-1]
    shape._inline.docPr.set("title", f"Figure {number}")
    shape._inline.docPr.set("descr", FIGURE_LEGENDS[number])
    caption = doc.add_paragraph(style="Caption")
    caption.paragraph_format.keep_together = True
    lead = caption.add_run(f"Figure {number}. ")
    lead.bold = True
    caption.add_run(FIGURE_LEGENDS[number])


def add_table(doc: Document, number: int) -> None:
    rows = load_table(number)
    if number == 1:
        caption_text = (
            "Full-precision external-test performance. F1 and MCC use the fixed 0.5 threshold."
        )
        widths = [900, 900, 2450, 950, 1000, 950, 950, 1260]
    else:
        caption_text = (
            "Paired ROC-AUC differences (A10 minus comparator) from 2,000 whole-protein bootstrap "
            "replicates."
        )
        widths = [800, 900, 2050, 1050, 1850, 900, 910, 900]
    assert sum(widths) == 9360

    cap = doc.add_paragraph(style="Caption")
    cap.paragraph_format.space_before = Pt(8)
    cap.paragraph_format.space_after = Pt(4)
    cap.paragraph_format.keep_with_next = True
    lead = cap.add_run(f"Table {number}. ")
    lead.bold = True
    cap.add_run(caption_text)

    tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.style = "Table Grid"
    apply_table_geometry(tbl, widths)
    set_repeat_table_header(tbl.rows[0])
    for table_row in tbl.rows:
        set_table_row_cant_split(table_row)
    for ridx, row in enumerate(rows):
        for cidx, value in enumerate(row):
            cell = tbl.cell(ridx, cidx)
            cell.text = value
            if ridx == 0:
                shade_cell(cell, DESIGN["tables"]["header_fill"])
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.space_after = Pt(1.5)
                para.paragraph_format.line_spacing = 1.0
                para.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                    if cidx in (0, 1, 2, 4)
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                for run in para.runs:
                    set_run_font(run, size=7.7 if number == 2 else 8.0,
                                 bold=(ridx == 0), color="111827")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def render_markdown() -> str:
    lines = [
        f"# {TITLE}",
        "",
        "**Article type:** Original Article  ",
        "**Target journal:** Bioinformatics Advances  ",
        f"**Running title:** {RUNNING_TITLE}",
        "",
        "**Authors:** [AUTHOR NAMES AND ORCID IDENTIFIERS]  ",
        "**Affiliations:** [AFFILIATIONS]  ",
        "**Correspondence:** [CORRESPONDING AUTHOR ADDRESS AND EMAIL]",
        "",
        "## Abstract",
        "",
        ABSTRACT,
        "",
        "**Keywords:** " + "; ".join(KEYWORDS),
        "",
    ]
    for item in CONTENT:
        kind = item["type"]
        if kind.startswith("h"):
            level = int(kind[1:]) + 1
            lines.extend(["#" * level + " " + item["text"], ""])
        elif kind == "p":
            lines.extend([item["text"], ""])
        elif kind == "bullets":
            lines.extend([f"- {entry}" for entry in item["items"]])
            lines.append("")
        elif kind == "figure":
            number = item["number"]
            lines.extend([
                f"![Figure {number}]({FIGURE_FILES[number].as_posix()})",
                "",
                f"**Figure {number}.** {FIGURE_LEGENDS[number]}",
                "",
            ])
        elif kind == "table":
            number = item["number"]
            table_path = C1 / (
                "table_1_external_performance.md"
                if number == 1
                else "table_2_paired_roc_statistics.md"
            )
            lines.extend([table_path.read_text(encoding="utf-8"), ""])
    lines.extend(["## References", ""])
    lines.extend([f"{ref}" for ref in REFERENCES])
    lines.append("")
    return "\n".join(lines)


def build_docx(path: Path) -> None:
    doc = Document()
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Locked external validation of intrinsic-disorder prediction"
    doc.core_properties.author = "Anonymous author group"
    doc.core_properties.last_modified_by = "Anonymous author group"
    doc.core_properties.keywords = ", ".join(KEYWORDS)
    configure_styles(doc)
    configure_page(doc)
    add_title_page(doc)
    add_abstract(doc)

    for item in CONTENT:
        kind = item["type"]
        if kind.startswith("h"):
            doc.add_heading(item["text"], level=int(kind[1:]))
        elif kind == "p":
            add_body_paragraph(doc, item["text"])
        elif kind == "bullets":
            add_bullet_list(doc, item["items"])
        elif kind == "figure":
            add_figure(doc, item["number"])
        elif kind == "table":
            add_table(doc, item["number"])

    doc.add_heading("References", level=1)
    for ref in REFERENCES:
        para = doc.add_paragraph(ref)
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.left_indent = Inches(0.25)
        para.paragraph_format.first_line_indent = Inches(-0.25)
        para.paragraph_format.space_after = Pt(5)
        para.paragraph_format.line_spacing = 1.1
    doc.save(path)


def build_cover_letter() -> str:
    return f"""# Cover letter draft

**To:** The Editors, *Bioinformatics Advances*
**Re:** {TITLE}

Dear Editors,

Please consider our manuscript, \"{TITLE},\" as an Original Article in the Sequence Analysis area. The work presents A10-locked, a frozen-ESM-2 ensemble designed and locked before CAID2/CAID3 label access. Its technical contribution combines multiscale residue context with a learned last-four-layer representation mixture, while its evaluation contribution is a hash-audited, homology-controlled protocol with paired whole-protein uncertainty across two CAID rounds and two operational disorder references.

The principal result is deliberately bounded: A10 performs strongly on the two Disorder-PDB tracks and long disordered segments, but it is not uniformly superior on Disorder-NOX. It is statistically comparable to a locally reproduced PUNCH2-Light Released-13 implementation across all four tracks. We believe this track-specific account, including negative results and reproducible baseline audits, will be useful to readers choosing or developing intrinsic-disorder predictors.

The manuscript is original, is not under consideration elsewhere, and all authors have approved submission [CONFIRM BEFORE USE]. CAID2/CAID3 labels were used only after model locking and never for training, tuning, recalibration or ensemble weighting. Source archives, checkpoints, predictions and reports are covered by SHA256 manifests; the public repository and DOI will be inserted before submission.

AI disclosure: OpenAI Codex assisted source organization, language drafting and programmatic Word assembly. Human authors are responsible for verifying every scientific claim, citation and final sentence, and AI was not used to select models or access evaluation labels during development.

Suggested editor/reviewer information and exclusions: [TO BE COMPLETED BY THE AUTHORS].

Sincerely,
[CORRESPONDING AUTHOR]
"""


def build_checklist() -> str:
    return f"""# C2 submission checklist

## Journal fit and limits

- Target: *Bioinformatics Advances*, Original Article, Sequence Analysis.
- Abstract: {word_count(ABSTRACT)} words (limit 200).
- Keywords: {len(KEYWORDS)} (limit 5).
- References: {len(REFERENCES)} (limit 50).
- Main figures: 3 (limit 6); supplementary figures: 2.
- Main tables: 2 (limit 6); supplementary tables: 6.
- Required sections present: Abstract, Introduction, Methods, Discussion, References.
- Word draft is an authoring layout, not the journal's final two-column template. Transfer to the official template and verify the 8-page limit before submission.

## Scientific integrity checks

- A10 selection and model locking preceded CAID2/CAID3 label access.
- Thirty checkpoint hashes were verified; the development checkpoint was excluded.
- CAID2/CAID3 labels were not used for training, tuning, weighting or recalibration.
- Protein-level paired bootstrap and Holm families are described exactly as executed.
- B7 subgroup and complementarity analyses are labeled descriptive.
- No claim of universal state-of-the-art performance is made.

## Author-owned fields still required

- [ ] Author names, order, affiliations and ORCID identifiers.
- [ ] Corresponding-author postal address and email.
- [ ] CRediT contribution statement.
- [ ] Funding and grant numbers.
- [ ] Conflict-of-interest statement.
- [ ] Acknowledgements.
- [ ] Repository URL/DOI and final software license.
- [ ] Permission review for third-party source code/weights; do not redistribute assets with unclear licenses.
- [ ] Suggested reviewers and exclusions.
- [ ] Human verification of every reference and all numerical claims.
- [ ] Cover-letter AI disclosure retained or adapted to journal policy.
"""


def copy_supplementary_assets() -> None:
    supp = OUT / "supplementary_materials"
    supp.mkdir(parents=True, exist_ok=True)
    assets = [
        "figure_4_disorder_fraction_heatmap.png",
        "figure_4_disorder_fraction_heatmap.svg",
        "figure_5_segment_error_profiles.png",
        "figure_5_segment_error_profiles.svg",
        "table_s1_all_paired_statistics.csv",
        "table_s2_all_subgroup_metrics.csv",
        "table_s3_all_segment_metrics.csv",
        "table_s4_all_boundary_metrics.csv",
        "table_s5_all_complementarity_metrics.csv",
        "table_s6_cross_dataset_robustness.csv",
    ]
    for name in assets:
        shutil.copy2(C1 / name, supp / name)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    RELEASE.mkdir(parents=True, exist_ok=True)
    for path in FIGURE_FILES.values():
        if not path.is_file():
            raise FileNotFoundError(path)

    abstract_words = word_count(ABSTRACT)
    if abstract_words > 200:
        raise ValueError(f"abstract has {abstract_words} words")
    if len(KEYWORDS) > 5 or len(REFERENCES) > 50:
        raise ValueError("journal limit exceeded")

    markdown_path = OUT / "A10_locked_manuscript.md"
    docx_path = OUT / "A10_locked_manuscript.docx"
    cover_path = OUT / "C2_COVER_LETTER_DRAFT.md"
    checklist_path = OUT / "C2_SUBMISSION_CHECKLIST.md"
    markdown_path.write_text(render_markdown(), encoding="utf-8", newline="\n")
    cover_path.write_text(build_cover_letter(), encoding="utf-8", newline="\n")
    checklist_path.write_text(build_checklist(), encoding="utf-8", newline="\n")
    build_docx(docx_path)
    copy_supplementary_assets()

    report = {
        "schema_version": 1,
        "status": "pass",
        "experiment": "c2_full_manuscript",
        "target_journal": "Bioinformatics Advances",
        "article_type": "Original Article",
        "title": TITLE,
        "abstract_words": abstract_words,
        "keywords": len(KEYWORDS),
        "references": len(REFERENCES),
        "main_figures": 3,
        "main_tables": 2,
        "design": DESIGN,
        "source_policy": "locked_aggregates_and_documented_protocol_only",
        "caid_labels_used_for_training_or_tuning": False,
        "outputs": {
            "docx": str(docx_path.relative_to(ROOT)).replace("\\", "/"),
            "docx_sha256": sha256_file(docx_path),
            "markdown": str(markdown_path.relative_to(ROOT)).replace("\\", "/"),
            "markdown_sha256": sha256_file(markdown_path),
            "cover_letter": str(cover_path.relative_to(ROOT)).replace("\\", "/"),
            "checklist": str(checklist_path.relative_to(ROOT)).replace("\\", "/"),
        },
        "unresolved_author_fields": True,
    }
    report_path = OUT / "c2_build_report.json"
    report_path.write_text(json.dumps(report, indent=2), encoding="utf-8", newline="\n")

    lock_path = OUT / "c2_result_lock.sha256"
    lock_targets = sorted(
        path for path in OUT.rglob("*")
        if (
            path.is_file()
            and path.name != lock_path.name
            and not any(part.startswith("_qa_render") for part in path.relative_to(OUT).parts)
        )
    )
    lock_lines = [
        f"{sha256_file(path)}  {path.relative_to(ROOT).as_posix()}"
        for path in lock_targets
    ]
    lock_path.write_text("\n".join(lock_lines) + "\n", encoding="utf-8", newline="\n")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
